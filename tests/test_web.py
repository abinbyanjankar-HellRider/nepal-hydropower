import io

import pytest
from openpyxl import load_workbook

from src import config
from src.collectors.manual_import import upsert_projects
from src.dashboard import create_app

RECORDS = [
    dict(project_id="HP_1", project_name_en="Alpha Khola", capacity_mw=100, status="Operational", district="Dolakha",
         province="Bagmati", latitude=27.6, longitude=86.2, commissioning_year=2020, developer="Alpha Power Ltd"),
    dict(project_id="HP_2", project_name_en="Beta <script>alert(1)</script>", capacity_mw=20,
         status="Under Construction", district="Ilam", province="Koshi", expected_completion_year=2027),
    dict(project_id="HP_3", project_name_en="Gamma", capacity_mw=5, status="Planned"),
]


@pytest.fixture()
def client(db, tmp_path, monkeypatch):
    monkeypatch.setattr(config, "EXPORT_DIR", tmp_path)
    with db.session_scope() as s:
        upsert_projects(s, RECORDS)
    app = create_app(db)
    app.config["TESTING"] = True
    return app.test_client()


def test_pages_render(client):
    for path in ("/", "/map", "/projects", "/projects/HP_1", "/companies", "/analytics"):
        r = client.get(path)
        assert r.status_code == 200, path
        assert b"Nepal Hydropower" in r.data


def test_projects_api_filters_sorting_and_paging(client):
    body = client.get("/api/projects").get_json()
    assert body["total"] == 3 and body["items"][0]["project_id"] == "HP_1"  # largest first
    assert client.get("/api/projects?status=Planned").get_json()["total"] == 1
    assert client.get("/api/projects?capacity_min=10&capacity_max=50").get_json()["items"][0]["project_id"] == "HP_2"
    assert client.get("/api/projects?q=alpha").get_json()["total"] == 1
    assert client.get("/api/projects?developer=alpha").get_json()["total"] == 1
    asc = client.get("/api/projects?sort=capacity_mw&order=asc&limit=1").get_json()
    assert asc["total"] == 3 and len(asc["items"]) == 1 and asc["items"][0]["project_id"] == "HP_3"


def test_api_rejects_bad_input_with_json_errors(client):
    for url in ("/api/projects?sort=nope", "/api/projects?capacity_min=abc", "/api/projects?limit=x",
                "/api/analytics/capacity?by=nope"):
        r = client.get(url)
        assert r.status_code == 400 and "error" in r.get_json(), url
    r = client.get("/api/projects/NOPE")
    assert r.status_code == 404 and r.get_json()["error"]


def test_non_integer_query_params_are_400_not_500(client):
    """Regression: top/days/limit were cast with a bare int(), so '?top=abc' crashed with a server error."""
    for url in ("/api/analytics/capacity?by=province&top=abc", "/api/analytics/licences?days=soon",
                "/api/analytics/unknowns?by=province&limit=x", "/api/news?limit=lots", "/api/projects?offset=1.5"):
        r = client.get(url)
        assert r.status_code == 400 and "must be an integer" in r.get_json()["error"], url
    assert client.get("/api/news?limit=0").status_code == 200  # out-of-range values are clamped, not rejected
    assert client.get("/api/analytics/capacity?by=province&top=").status_code == 200  # empty -> default


def test_next_project_id_compares_numbers_not_text():
    from src.dashboard import next_project_id
    assert next_project_id(["HP_1", "HP_2", "HP_3"]) == "HP_004"
    assert next_project_id(["HP_998", "HP_999", "HP_1000"]) == "HP_1001"  # text order would pick HP_999
    assert next_project_id(["X_9", None]) == "HP_001"


def test_news_shows_listed_company_name(db, client):
    """Regression: news showed the DoED promoter name ('Himal Hydro') instead of the NEPSE listed name."""
    from datetime import date
    from src.models import Company, ProjectUpdate
    with db.session_scope() as s:
        c = Company(company_name="Himal Hydro", listed_name="Super Madi Hydropower Limited", nepse_listed=True, stock_symbol="SMHL")
        s.add(c)
        s.flush()
        s.add(ProjectUpdate(company_id=c.company_id, update_date=date(2026, 9, 1), title="SMHL results",
                            source_name="Test", source_url="https://example.org/a"))
    item = client.get("/api/news?scope=companies").get_json()[0]
    assert item["company"] == "Super Madi Hydropower Limited" and item["company_symbol"] == "SMHL"


def test_json_has_no_nan(client):
    raw = client.get("/api/projects").get_data(as_text=True)
    assert "NaN" not in raw  # NaN is invalid JSON; missing values must be null


def test_project_detail_statistics_and_geojson(client):
    p = client.get("/api/projects/HP_1").get_json()
    assert p["developer"] == "Alpha Power Ltd" and p["district"] == "Dolakha" and p["financials"] == []
    stats = client.get("/api/projects/statistics").get_json()
    assert stats["summary"]["operational_mw"] == 100
    gj = client.get("/api/geojson").get_json()
    assert [f["properties"]["project_id"] for f in gj["features"]] == ["HP_1"]  # only rows with coordinates
    assert gj["features"][0]["geometry"]["coordinates"] == [86.2, 27.6]


def test_hostile_names_are_returned_as_data_not_markup(client):
    body = client.get("/api/projects?q=beta").get_json()
    assert body["items"][0]["name"] == "Beta <script>alert(1)</script>"
    # pages never inline project data into HTML; it only travels through JSON + textContent
    assert b"alert(1)" not in client.get("/projects/HP_2").data
    assert b"alert(1)" not in client.get("/projects").data


def test_admin_add_requires_token(client, monkeypatch):
    payload = {"project_name_en": "New Khola", "capacity_mw": 12, "status": "Planned"}
    monkeypatch.delenv("ADMIN_TOKEN", raising=False)
    assert client.post("/api/projects/add", json=payload).status_code == 403  # disabled by default
    monkeypatch.setenv("ADMIN_TOKEN", "s3cret")
    assert client.post("/api/projects/add", json=payload, headers={"X-Admin-Token": "wrong"}).status_code == 403
    ok = client.post("/api/projects/add", json=payload, headers={"X-Admin-Token": "s3cret"})
    assert ok.status_code == 201 and ok.get_json()["project_id"] == "HP_004"
    bad = client.post("/api/projects/add", json={"project_name_en": "", "capacity_mw": 99999},
                      headers={"X-Admin-Token": "s3cret"})
    assert bad.status_code == 400 and len(bad.get_json()["errors"]) == 2
    assert client.get("/api/projects").get_json()["total"] == 4


def test_excel_export_endpoint(client):
    r = client.get("/api/export/excel?format=master")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.data))
    assert {"Summary", "Projects", "Technical", "Ownership", "Timeline", "Grid", "Financial"} <= set(wb.sheetnames)
    assert wb["Projects"].max_row == 4  # header + 3 projects
    assert wb["Summary"]["B5"].value.startswith("=COUNTIF(Projects!")  # live formulas, not pasted numbers
    assert client.get("/api/export/excel?format=bogus").status_code == 400


def test_excel_export_never_turns_data_into_formulas():
    """Regression (QA): a scraped name such as '=HYPERLINK(...)' was written as a live formula in the data sheets."""
    import pandas as pd
    from openpyxl import Workbook
    from src.exports.excel_builder import _write_sheet
    wb = Workbook()
    _write_sheet(wb, "Data", pd.DataFrame({"name": ["=HYPERLINK(\"http://evil\",\"x\")", "@SUM(1)", "+cmd", "Alpha"], "mw": [1.0, 2.0, 3.0, 4.0]}))
    buf = io.BytesIO()
    wb.save(buf)
    ws = load_workbook(io.BytesIO(buf.getvalue()))["Data"]
    assert ws["A2"].data_type == "s" and ws["A2"].value.startswith("=HYPERLINK")     # text, not a formula
    assert ws["A3"].data_type == "s" and ws["A4"].data_type == "s" and ws["A5"].value == "Alpha"
    assert ws["B2"].value == 1.0                                                      # numbers are untouched


def test_json_never_contains_nan_or_infinity(db, client):
    """Regression: a listed company with zero paid-up capital produced `Infinity`, which broke the Companies page."""
    from src.collectors.nepse_scraper import sync_listed_companies
    with db.session_scope() as s:
        sync_listed_companies(s, [{"symbol": "ZERO", "name": "Alpha Power Ltd", "listed_shares": 0, "paidup_value": 0.0}],
                              aliases={})
    for url in ("/api/companies", "/api/projects", "/api/projects/statistics", "/api/geojson", "/api/projects/HP_1"):
        text = client.get(url).get_data(as_text=True)
        assert "Infinity" not in text and "NaN" not in text, url
    row = next(r for r in client.get("/api/companies").get_json()["nepse"] if r["symbol"] == "ZERO")
    assert row["op_mw_per_bn_paidup"] is None
