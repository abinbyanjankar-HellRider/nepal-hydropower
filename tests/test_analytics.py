from datetime import date

import pandas as pd
import pytest

from src.analytics import financials as fin
from src.analytics import load_projects_df
from src.analytics import ownership as own
from src.analytics import technical as tech
from src.collectors.manual_import import import_financials_csv, upsert_projects
from src.collectors.nepse_scraper import sync_listed_companies

RECORDS = [
    dict(project_id="HP_1", project_name_en="Alpha", capacity_mw=100, status="Operational", district="Dolakha",
         commissioning_year=2020, developer="Alpha Power Company Limited", annual_energy_generation_gwh=500),
    dict(project_id="HP_2", project_name_en="Beta", capacity_mw=50, status="Operational", district="Gorkha",
         commissioning_year=2020, developer="Alpha Power Company Ltd."),
    dict(project_id="HP_3", project_name_en="Gamma", capacity_mw=25, status="Operational", district="Ilam",
         commissioning_year=2018, developer="Gamma Hydro Ltd"),
    dict(project_id="HP_4", project_name_en="Delta", capacity_mw=200, status="Under Construction", district="Dolakha",
         developer="Alpha Power Company Limited", expected_completion_year=2027,
         license_expiry_date="2020-01-01", license_type="Generation"),
    dict(project_id="HP_5", project_name_en="Eps", capacity_mw=0.5, status="Planned"),
]


@pytest.fixture()
def df(db):
    with db.session_scope() as s:
        upsert_projects(s, RECORDS)
        sync_listed_companies(s, [{"symbol": "ALPHA", "name": "Alpha Power Co. Ltd.", "listed_shares": 10_000_000,
                                   "paidup_value": 100.0}])
    with db.session_scope() as s:
        return load_projects_df(s)


def test_capacity_by_status_and_shares(df):
    out = tech.capacity_by(df, "status").set_index("group")
    assert out.loc["Operational", "total_mw"] == 175 and out.loc["Operational", "projects"] == 3
    assert out["share_pct"].sum() == pytest.approx(100, abs=0.2)


def test_capacity_by_size_buckets(df):
    out = tech.capacity_by(df, "size").set_index("group")
    assert out.loc["<1 MW", "projects"] == 1 and out.loc["100-500 MW", "projects"] == 2


def test_commissioning_timeline_is_cumulative(df):
    t = tech.commissioning_timeline(df)
    assert t["year"].tolist() == [2018, 2020]
    assert t["cumulative_mw"].tolist() == [25, 175]


def test_pipeline_and_licence_watch(df):
    assert tech.pipeline_by_year(df).set_index("expected_year").loc[2027, "total_mw"] == 200
    watch = tech.expiring_licences(df, today=date(2026, 1, 1))
    assert watch["project_id"].tolist() == ["HP_4"] and watch["days_left"].iloc[0] < 0


def test_company_name_variants_share_one_company_and_nepse_exposure(df):
    ranking = own.developer_ranking(df).set_index("developer")
    assert ranking.loc["Alpha Power Company Limited", "op_projects"] == 2  # 'Ltd.' variant folded in
    assert ranking.loc["Alpha Power Company Limited", "pipe_mw"] == 200


def test_nepse_exposure_uses_listing_data(db, df):
    with db.session_scope() as s:
        exp = own.nepse_exposure(df, s).set_index("symbol")
    row = exp.loc["ALPHA"]
    assert row["op_mw"] == 150 and row["pipe_mw"] == 200 and row["projects"] == 3
    assert row["paidup_npr_m"] == 1000  # 10M shares x NPR 100


def test_company_lookup_by_symbol_and_fragment(df):
    assert len(own.company_projects(df, "ALPHA")) == 3
    assert len(own.company_projects(df, "gamma")) == 1


def test_revenue_estimate_uses_project_energy_when_known(df):
    est = fin.estimate_revenue_df(df, capacity_factor=0.5, wet_tariff=5, dry_tariff=5, dry_share=0.3)
    by = est.set_index("project_id")
    assert by.loc["HP_1", "est_energy_gwh"] == 500 and "project energy" in by.loc["HP_1", "basis"]
    assert by.loc["HP_2", "est_energy_gwh"] == pytest.approx(50 * 8.76 * 0.5, rel=1e-3)
    assert by.loc["HP_2", "est_revenue_npr_m"] == pytest.approx(by.loc["HP_2", "est_energy_gwh"] * 5, rel=1e-3)


def _plants():
    return pd.DataFrame([
        dict(project_id="P1", name="One", status="Operational", capacity_mw=10.0, annual_energy_generation_gwh=None,
             ppa_rate_npr_per_kwh=None, developer="A", symbol="A", company_id=1),
        dict(project_id="P2", name="Two", status="Operational", capacity_mw=10.0, annual_energy_generation_gwh=None,
             ppa_rate_npr_per_kwh=None, developer="B", symbol="B", company_id=2),
        dict(project_id="P3", name="Three", status="Operational", capacity_mw=10.0, annual_energy_generation_gwh=None,
             ppa_rate_npr_per_kwh=None, developer="C", symbol="C", company_id=3),
        dict(project_id="P4", name="Four", status="Operational", capacity_mw=10.0, annual_energy_generation_gwh=None,
             ppa_rate_npr_per_kwh=5.5, developer="D", symbol="D", company_id=4)])


def test_revenue_estimate_prices_plants_at_verified_ppa_rates_and_says_which(df):
    """QA: projects.ppa_rate_npr_per_kwh is empty, so every plant was priced at the assumed tariff and ignored the
    verified PPA rates that the /income page uses."""
    rates = {(1, "P1"): (4.0, 7.0), (2, None): (4.8, 8.4)}
    est = fin.estimate_revenue_df(_plants(), capacity_factor=0.5, dry_share=0.3, rates=rates).set_index("project_id")
    assert est.loc["P1", "tariff_npr_kwh"] == pytest.approx(4.0 * 0.7 + 7.0 * 0.3) and est.loc["P1", "tariff_basis"] == "verified project"
    assert est.loc["P2", "tariff_npr_kwh"] == pytest.approx(4.8 * 0.7 + 8.4 * 0.3) and est.loc["P2", "tariff_basis"] == "verified company"
    assert est.loc["P3", "tariff_basis"] == "assumed"                      # no verified rate for company 3
    assert est.loc["P4", "tariff_npr_kwh"] == 5.5 and est.loc["P4", "tariff_basis"] == "project record"


def test_revenue_estimate_without_rates_is_all_assumed_as_before():
    est = fin.estimate_revenue_df(_plants(), capacity_factor=0.5, dry_share=0.3)
    assert set(est["tariff_basis"]) == {"assumed", "project record"}


def test_excel_links_the_assumption_cell_only_to_truly_assumed_tariffs():
    """QA: a plant whose verified tariff happened to equal the assumed blend (4.8/8.4) was tied to the assumption cell."""
    from openpyxl import Workbook
    from src.exports.excel_builder import _link_financial_formulas
    est = fin.estimate_revenue_df(_plants().iloc[1:3], rates={(2, None): (4.8, 8.4)})     # P2 verified 5.88, P3 assumed 5.88
    assert est.loc[est["project_id"] == "P2", "tariff_npr_kwh"].iloc[0] == est.loc[est["project_id"] == "P3", "tariff_npr_kwh"].iloc[0]
    ws = Workbook().active
    _link_financial_formulas(ws, est)
    assert ws.cell(4, 5).value is None                          # P2: verified, stays a value (set by the data sheet)
    assert ws.cell(5, 5).value == "=Assumptions!$B$7"           # P3: truly assumed, live link


def test_projects_frame_carries_the_company_id(df):
    assert "company_id" in df.columns and df["company_id"].notna().any()


def test_irr_known_values():
    assert fin.irr([-100, 110]) == pytest.approx(0.10, abs=1e-6)
    assert fin.irr([-1000, 300, 300, 300, 300, 300]) == pytest.approx(0.1524, abs=1e-3)
    assert fin.irr([100, 100]) is None  # no outlay: undefined
    profile = fin.project_return_profile(1000, 100, life_years=35)
    assert profile["simple_payback_years"] == 10.0 and 9 < profile["irr_pct"] < 10


def test_financial_history_derives_capacity_factor(db, tmp_path, df):
    f = tmp_path / "f.csv"
    f.write_text("project_id,fiscal_year,energy_generated_gwh,revenue_npr\nHP_1,2080/81,438,2500000000\n")
    with db.session_scope() as s:
        import_financials_csv(s, f)
        h = fin.financial_history(s, "HP_1")
        port = fin.portfolio_financials(s)
    assert h["capacity_factor_pct"].iloc[0] == pytest.approx(50.0, abs=0.1)  # 438 / (100 MW x 8.76)
    assert port["projects"].iloc[0] == 1


def test_forecast_capacity_scenario(df):
    base = tech.forecast_capacity(df, start_year=2026, horizon=3)
    assert base["cumulative_mw"].iloc[0] == 175  # operational only until Delta lands
    assert base.set_index("year").loc[2027, "mw_added"] == 200 and base["cumulative_mw"].iloc[-1] == 375
    slipped = tech.forecast_capacity(df, delay_years=1, start_year=2026, horizon=3, completion_rate=0.5)
    assert slipped.set_index("year").loc[2028, "mw_added"] == 100  # 200 MW x 50%, one year later
    # a past expected year lands in the first forecast year instead of vanishing
    assert tech.forecast_capacity(df, start_year=2030, horizon=1)["mw_added"].iloc[0] == 200


def test_nepse_matching_is_strict_and_never_shares_a_company(db):
    """Regression: 'Super Madi' (SMHL) was fuzzy-linked to 'Supermai', stealing SMH's company."""
    from sqlalchemy import select
    from src.models import Company
    with db.session_scope() as s:
        upsert_projects(s, [{"project_id": "HP_1", "project_name_en": "X", "developer": "Supermai Hydropower Pvt.Ltd."}])
        res = sync_listed_companies(s, [
            {"symbol": "SMH", "name": "Super Mai Hydropower Limited", "listed_shares": 1, "paidup_value": 100.0},
            {"symbol": "SMHL", "name": "Super Madi Hydropower Limited", "listed_shares": 2, "paidup_value": 100.0}])
        listed = {c.stock_symbol: c.company_name for c in s.scalars(select(Company).where(Company.nepse_listed.is_(True)))}
    assert set(listed) == {"SMH", "SMHL"}  # every symbol keeps its own company row
    assert listed["SMH"] == "Supermai Hydropower Pvt.Ltd."  # matched despite spacing difference
    assert listed["SMHL"] == "Super Madi Hydropower Limited"  # NOT linked to the Supermai promoter
    assert res == {"matched": 1, "created": 1}


def test_nepse_alias_links_a_confirmed_promoter_and_resync_resets_stale_flags(db):
    from sqlalchemy import select
    from src.models import Company
    with db.session_scope() as s:
        upsert_projects(s, [{"project_id": "HP_1", "project_name_en": "X", "developer": "Mathillo Mailung Khola Jalvidyut Ltd."}])
        row = {"symbol": "MMKJL", "name": "Mathillo Mailun Khola Jalvidhyut Limited", "listed_shares": 1, "paidup_value": 100.0}
        sync_listed_companies(s, [row], aliases={"MMKJL": "Mathillo Mailung Khola Jalvidyut Ltd."})
        linked = s.scalars(select(Company).where(Company.stock_symbol == "MMKJL")).one()
        assert linked.company_name == "Mathillo Mailung Khola Jalvidyut Ltd."
        assert linked.listed_name == "Mathillo Mailun Khola Jalvidhyut Limited"  # NEPSE name kept alongside DoED's
        sync_listed_companies(s, [], aliases={})  # company delisted: flags must be cleared
        assert s.scalars(select(Company).where(Company.nepse_listed.is_(True))).all() == []
        assert linked.listed_name is None


def test_profile_shows_listed_name_for_aliased_promoter(db):
    """Regression: SMHL showed as 'Himal Hydro' (its DoED promoter name) instead of Super Madi Hydropower."""
    from sqlalchemy import select
    from src.models import Company
    from src.analytics.company_profiles import build_profile
    with db.session_scope() as s:
        upsert_projects(s, [{"project_id": "HP_1", "project_name_en": "Super Madi", "developer": "Himal Hydro"}])
        sync_listed_companies(s, [{"symbol": "SMHL", "name": "Super Madi Hydropower Limited", "listed_shares": 1,
                                   "paidup_value": 100.0}], aliases={"SMHL": "Himal Hydro"})
        c = s.scalars(select(Company).where(Company.stock_symbol == "SMHL")).one()
        prof = build_profile(s, c.company_id)["company"]
    assert prof["name"] == "Super Madi Hydropower Limited"
    assert prof["doed_name"] == "Himal Hydro"
