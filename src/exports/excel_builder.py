"""Excel report generation (openpyxl). Summary sheets use live formulas over the data sheets."""
from __future__ import annotations

from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from .. import config
from ..analytics import financials as fin
from ..analytics import income
from ..analytics import load_projects_df
from ..analytics import ownership as own
from ..analytics import technical as tech
from ..database import DatabaseManager, db_manager as default_db
from ..models import Project, ProjectFinancial

KINDS = ("master", "operational", "construction", "planned", "financial", "technical", "nepse")
STATUS_FILTER = {"operational": "Operational", "construction": "Under Construction", "planned": "Planned"}

HEADER_FILL = PatternFill("solid", fgColor="0B4F6C")
HEADER_FONT = Font(bold=True, color="FFFFFF")
INPUT_FONT = Font(color="0000FF")  # blue = hard-coded assumption the user may change
NUM_FORMATS = {"mw": "#,##0.0", "int": "#,##0", "money": "#,##0", "pct": "0.0", "coord": "0.0000"}

PROJECT_COLS = ["project_id", "name", "status", "capacity_mw", "project_type", "province", "district", "river_name_raw",
                "developer", "symbol", "commissioning_year", "expected_completion_year", "license_type",
                "license_number", "license_expiry_date", "latitude", "longitude", "data_reliability"]


def _full_frame(db: DatabaseManager) -> pd.DataFrame:
    """All Project columns (enums flattened) plus developer/owner names."""
    with db.session_scope() as s:
        rows = []
        for p in s.scalars(select(Project)):
            row = {}
            for col in p.__table__.columns:
                v = getattr(p, col.name)
                row[col.name] = getattr(v, "value", v)
            row["developer"] = p.developer.company_name if p.developer else None
            row["owner"] = p.owner.company_name if p.owner else None
            row["district"] = p.district.district_name if p.district else None
            rows.append(row)
    return pd.DataFrame(rows).rename(columns={"project_name_en": "name"})


def _write_sheet(wb: Workbook, title: str, df: pd.DataFrame, formats: dict[str, str] | None = None,
                 note: str | None = None) -> None:
    ws = wb.create_sheet(title[:31])
    start = 1
    if note:
        ws.cell(1, 1, note).font = Font(italic=True, color="7A5200")
        start = 3
    for j, col in enumerate(df.columns, 1):
        c = ws.cell(start, j, col)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, row in enumerate(df.itertuples(index=False), start + 1):
        for j, v in enumerate(row, 1):
            if v is None or (isinstance(v, float) and pd.isna(v)) or v is pd.NaT:
                continue
            if hasattr(v, "item"):
                v = v.item()
            cell = ws.cell(i, j, v)
            if isinstance(v, str) and v.startswith("="):
                cell.data_type = "s"  # scraped text is data, never a formula (openpyxl would store '=...' as one)
    for j, col in enumerate(df.columns, 1):
        fmt = (formats or {}).get(col)
        if fmt:
            for i in range(start + 1, start + 1 + len(df)):
                ws.cell(i, j).number_format = fmt
        width = max([len(str(col))] + [len(str(v)) for v in df[col].head(200).tolist() if v is not None]) + 2
        ws.column_dimensions[get_column_letter(j)].width = min(max(width, 8), 46)
    ws.freeze_panes = ws.cell(start + 1, 1)
    if len(df):
        ws.auto_filter.ref = f"A{start}:{get_column_letter(len(df.columns))}{start + len(df)}"


def _summary_sheet(wb: Workbook, n_projects: int, projects_cols: list[str], generated: date) -> None:
    """Counts and MW by stage as live COUNTIF/SUMIF formulas over the 'Projects' sheet."""
    ws = wb.create_sheet("Summary", 0)
    ws["A1"], ws["A1"].font = "Nepal hydropower: summary", Font(bold=True, size=14)
    ws["A2"] = (f"Generated {generated:%Y-%m-%d} from DoED / Wikipedia / NEPSE listings. "
                "Figures below are formulas over the 'Projects' sheet.")
    for j, h in enumerate(["Stage", "Projects", "Capacity (MW)", "Share of MW"], 1):
        c = ws.cell(4, j, h)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    sc, cc = get_column_letter(projects_cols.index("status") + 1), get_column_letter(projects_cols.index("capacity_mw") + 1)
    last = n_projects + 1
    stages = ["Operational", "Under Construction", "Licensed", "Planned"]
    for i, stage in enumerate(stages, 5):
        ws.cell(i, 1, stage)
        ws.cell(i, 2, f'=COUNTIF(Projects!${sc}$2:${sc}${last},A{i})')
        ws.cell(i, 3, f'=SUMIF(Projects!${sc}$2:${sc}${last},A{i},Projects!${cc}$2:${cc}${last})')
        ws.cell(i, 4, f"=IF($C${5 + len(stages)}=0,0,C{i}/$C${5 + len(stages)})")
        ws.cell(i, 3).number_format, ws.cell(i, 4).number_format = NUM_FORMATS["mw"], "0.0%"
    t = 5 + len(stages)
    ws.cell(t, 1, "Total").font = Font(bold=True)
    ws.cell(t, 2, f"=SUM(B5:B{t - 1})")
    ws.cell(t, 3, f"=SUM(C5:C{t - 1})")
    ws.cell(t, 3).number_format = NUM_FORMATS["mw"]
    for col in "ABCD":
        ws.cell(t, "ABCD".index(col) + 1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 28
    for col in "BCD":
        ws.column_dimensions[col].width = 16


def _assumptions_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Assumptions")
    ws["A1"], ws["A1"].font = "Revenue-estimate assumptions (blue = editable input)", Font(bold=True, size=12)
    rows = [("Wet-season tariff (NPR/kWh)", config.ASSUMED_WET_TARIFF_NPR_PER_KWH),
            ("Dry-season tariff (NPR/kWh)", config.ASSUMED_DRY_TARIFF_NPR_PER_KWH),
            ("Dry-season share of annual energy", config.ASSUMED_DRY_SEASON_ENERGY_SHARE),
            ("Assumed capacity factor", config.ASSUMED_CAPACITY_FACTOR)]
    for i, (label, value) in enumerate(rows, 3):
        ws.cell(i, 1, label)
        c = ws.cell(i, 2, value)
        c.font = INPUT_FONT
    ws.cell(7, 1, "Blended tariff (NPR/kWh)")
    ws.cell(7, 2, "=B3*(1-B5)+B4*B5")
    ws.cell(9, 1, "These are assumptions, not sourced data: replace with the current NEA PPA rate schedule. "
                  "The Financial sheet's revenue column references B7 and B6.").font = Font(italic=True, color="7A5200")
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 14


def _link_financial_formulas(ws, est: pd.DataFrame) -> None:
    """Make the estimate columns live: change the Assumptions sheet and revenue recalculates.

    Columns: A id, B name, C MW, D energy GWh, E tariff, F revenue (NPR m), G basis. Header is on row 3.
    Only cells that were derived from an assumption are linked; project-reported energy / PPA rates stay values.
    """
    for i, row in enumerate(est.itertuples(index=False), 4):
        if str(row.basis).startswith("assumed"):
            ws.cell(i, 4, f"=C{i}*8.76*Assumptions!$B$6")
        if row.tariff_basis == "assumed":  # never tie a verified or recorded rate to the assumption just because it is equal
            ws.cell(i, 5, "=Assumptions!$B$7")
        ws.cell(i, 6, f"=D{i}*E{i}")
        for col in (4, 5, 6):
            ws.cell(i, col).number_format = "#,##0.00"


def build_workbook(kind: str = "master", db: DatabaseManager | None = None, out_dir: Path | None = None) -> Path:
    if kind not in KINDS:
        raise ValueError(f"format must be one of {list(KINDS)}")
    db = db or default_db
    out_dir = Path(out_dir or config.EXPORT_DIR)
    out_dir.mkdir(parents=True, exist_ok=True)
    with db.session_scope() as s:
        flat = load_projects_df(s)
        ppa_rates = income.load_ppa_rates(s)  # verified PPA rates price the revenue estimate, as on the /income page
        nepse = own.nepse_exposure(flat, s) if not flat.empty else pd.DataFrame()
        reported = pd.DataFrame([{
            "project_id": f.project_id, "fiscal_year": f.fiscal_year, "energy_gwh": f.energy_generated_gwh,
            "revenue_npr": f.revenue_npr, "operating_cost_npr": f.operating_cost_npr,
            "net_profit_npr": f.net_profit_npr, "capacity_factor_pct": f.capacity_factor_percentage,
            "irr_pct": f.irr_percentage, "source": f.source}
            for f in s.scalars(select(ProjectFinancial))])
    full = _full_frame(db)
    if flat.empty:
        raise ValueError("database has no projects; run `python main.py sync` first")
    full = full.merge(flat[["project_id", "symbol"]], on="project_id", how="left")

    wb = Workbook()
    wb.remove(wb.active)
    name_of = {"status": "status", "capacity_mw": "capacity_mw"}
    mw = {"capacity_mw": NUM_FORMATS["mw"], "latitude": NUM_FORMATS["coord"], "longitude": NUM_FORMATS["coord"]}

    if kind in ("master", "operational", "construction", "planned"):
        d = full[PROJECT_COLS].sort_values("capacity_mw", ascending=False)
        if kind in STATUS_FILTER:
            d = d[d["status"] == STATUS_FILTER[kind]]
        _write_sheet(wb, "Projects", d, mw)
    if kind == "master":
        _summary_sheet(wb, len(full), PROJECT_COLS, date.today())
        tech_cols = ["project_id", "name", "status", "capacity_mw", "project_type", "turbine_type", "turbine_count",
                     "design_head_m", "design_discharge_m3s", "tunnel_length_m", "penstock_length_m",
                     "annual_energy_generation_gwh"]
        _write_sheet(wb, "Technical", full[tech_cols], {"capacity_mw": NUM_FORMATS["mw"]})
        own_cols = ["project_id", "name", "status", "capacity_mw", "developer", "owner", "symbol"]
        _write_sheet(wb, "Ownership", full[own_cols], mw)
        time_cols = ["project_id", "name", "status", "construction_start_date", "commissioning_date",
                     "commissioning_year", "expected_completion_year", "license_issue_date", "license_expiry_date"]
        _write_sheet(wb, "Timeline", full[time_cols])
        grid_cols = ["project_id", "name", "status", "transmission_voltage_kv", "transmission_length_km",
                     "substation_name", "grid_connected"]
        _write_sheet(wb, "Grid", full[grid_cols])
    if kind == "technical":
        cols = ["project_id", "name", "status", "capacity_mw", "project_type", "turbine_type", "turbine_count",
                "design_head_m", "design_discharge_m3s", "tunnel_length_m", "penstock_length_m",
                "annual_energy_generation_gwh", "transmission_voltage_kv"]
        _write_sheet(wb, "Technical", full[cols].sort_values("capacity_mw", ascending=False),
                     {"capacity_mw": NUM_FORMATS["mw"]})
        _write_sheet(wb, "Capacity by province", tech.capacity_by(flat, "province"), {"total_mw": NUM_FORMATS["mw"]})
        _write_sheet(wb, "Commissioning", tech.commissioning_timeline(flat), {"mw_added": NUM_FORMATS["mw"],
                                                                              "cumulative_mw": NUM_FORMATS["mw"]})
    if kind in ("master", "financial"):
        est = fin.estimate_revenue_df(flat, rates=ppa_rates)
        _write_sheet(wb, "Financial", est.rename(columns={"est_revenue_npr_m": "est_revenue_npr_m (see Assumptions)"}),
                     {"capacity_mw": NUM_FORMATS["mw"], "est_energy_gwh": NUM_FORMATS["mw"]},
                     note="ESTIMATES for operational plants (project energy where known, else assumed capacity factor; "
                          "tariff_basis says whether a verified PPA rate or the assumed tariff was used). Not reported revenue.")
        _assumptions_sheet(wb)
        _link_financial_formulas(wb["Financial"], est)
        if not reported.empty:
            _write_sheet(wb, "Reported financials", reported, {"revenue_npr": NUM_FORMATS["money"]})
    if kind in ("master", "nepse"):
        _write_sheet(wb, "NEPSE exposure", nepse, {"op_mw": NUM_FORMATS["mw"], "pipe_mw": NUM_FORMATS["mw"],
                                                   "planned_mw": NUM_FORMATS["mw"]},
                     note="Direct holdings of record from DoED promoter names; excludes subsidiaries and minority stakes.")
    path = out_dir / f"hydropower_{kind}.xlsx"
    wb.save(path)
    return path
